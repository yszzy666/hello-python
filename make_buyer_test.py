import openpyxl
from openpyxl.styles import Font, Alignment

# 这份是模拟"买家真的发过来"的表：长相和上一份不一样
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "高一三班"

# 第 1 行整行空着
# 第 2 行是合并标题
ws.merge_cells("A2:D2")
ws["A2"] = "期中作业提交登记表"
ws["A2"].font = Font(bold=True, size=14)
ws["A2"].alignment = Alignment(horizontal="center")

# 第 3 行才是表头，而且列名和上次不一样
表头 = ["序号", "学生姓名", "提交情况", "备注"]
for i, h in enumerate(表头, start=1):
    ws.cell(row=3, column=i, value=h).font = Font(bold=True)

写法 = ["是", "已交", "交了", "未交", "否", "√", "1", "0", "请假", "", "下周三交", "是"]

姓名 = [
    "王小二", "李大明", "张小雨", "刘子轩", "陈思远", "杨浩然", "黄嘉怡", "周文博",
    "吴晓峰", "徐若曦", "孙志强", "马佳宁", "朱天佑", "胡一鸣", "郭子涵", "林诗雨",
    "何俊杰", "高梦琪", "罗宇航", "郑欣怡", "梁家豪", "谢雨欣", "唐嘉伟",
]

for i, name in enumerate(姓名, start=1):
    r = i + 3
    ws.cell(row=r, column=1, value=i)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=写法[(i - 1) % len(写法)])
    if name == "张小雨":
        ws.cell(row=r, column=4, value="请假，下周补交")
    if name == "胡一鸣":
        ws.cell(row=r, column=4, value="已私聊说明")

# 中间空一行（真实表格经常有）
# 最后一行是合计——这个最容易被当成一个人
r = len(姓名) + 5
ws.cell(row=r, column=1, value="")
ws.cell(row=r, column=2, value="合计")
ws.cell(row=r, column=3, value="23人")

ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 18

wb.save("data_buyer.xlsx")
print("已生成 data_buyer.xlsx")
print("这份表和上一份不一样：")
print("  · 第 1 行整行是空的，第 2 行是合并标题，第 3 行才是表头")
print("  · 列名是「学生姓名」「提交情况」，不是「姓名」「是否提交」")
print("  · 多了一列「备注」")
print("  · 最后一行是「合计 23人」")
