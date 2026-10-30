import openpyxl
from openpyxl.styles import Font, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "作业"

# ===== 第 1 行：班级标题（合并单元格，长这样是真实班级表的常见样子）=====
ws.merge_cells("A1:D1")
ws["A1"] = "高一(3)班 语文作业提交情况"
ws["A1"].font = Font(bold=True, size=14)
ws["A1"].alignment = Alignment(horizontal="center")

# ===== 第 2 行：真正的表头（注意：不在第 1 行）=====
表头 = ["序号", "姓名", "学号", "提交情况"]
for i, h in enumerate(表头, start=1):
    ws.cell(row=2, column=i, value=h).font = Font(bold=True)

# ===== 第 3 行起：40 个人，混着各种写法 =====
写法 = [
    "是", "已交", "交了", "提交", "√", "1", " 是 ", "是  ",
    "否", "没交", "未交", "0", "请假", "？", "", "下周三交",
    "已提交", "完成", "x", "×",
]

姓名表 = [
    "张三", "李四", "王五", "赵六", "钱七", "孙八", "周九", "吴十",
    "郑十一", "王十二", "冯十三", "陈十四", "褚十五", "卫十六", "蒋十七", "沈十八",
    "韩十九", "杨二十", "朱廿一", "秦廿二", "尤廿三", "许廿四", "何廿五", "吕廿六",
    "施廿七", "张廿八", "孔廿九", "曹三十", "严卅一", "华卅二", "金卅三", "魏卅四",
    "陶卅五", "姜卅六", "戚卅七", "谢卅八", "邹卅九", "喻四十", "柏四一", "水四二",
]

for i, name in enumerate(姓名表, start=1):
    r = i + 2
    ws.cell(row=r, column=1, value=i)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=2026000 + i)
    ws.cell(row=r, column=4, value=写法[(i - 1) % len(写法)])

# ===== 再补一行重名（真实名单里同名同姓很常见）=====
ws.cell(row=43, column=1, value=41)
ws.cell(row=43, column=2, value="张三")
ws.cell(row=43, column=3, value=2026001)
ws.cell(row=43, column=4, value="否")

ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 14

wb.save("data_real.xlsx")
print("已生成 data_real.xlsx —— 41 行数据，表头在第 2 行，第 1 行是班级标题")
print("里面混了 20 种不同的提交写法，还有 2 个空格子、1 个重名（张三）")
